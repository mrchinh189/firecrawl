import pathlib
import zipfile
from docx import Document
from report import build_docx


def test_build_docx_structure(tmp_path):
    out = tmp_path / "bao_cao.docx"
    path = build_docx.build(out_path=str(out))
    assert pathlib.Path(path).exists()
    doc = Document(path)
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    full = "\n".join(p.text for p in doc.paragraphs)
    for kw in ["Tổng quan", "Phân tích", "GIÁ CHUNG", "Cảnh báo", "Dự báo", "Nguồn", "Tin tức"]:
        assert any(kw in h for h in heads), f"thiếu heading chứa: {kw}"
    assert len(doc.tables) >= 2
    assert "2026-06-" in full


def test_build_docx_has_hyperlink(tmp_path):
    out = tmp_path / "b.docx"
    path = build_docx.build(out_path=str(out))
    with zipfile.ZipFile(path) as z:
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8", "ignore")
    assert "http" in rels  # có ít nhất 1 hyperlink ra ngoài
