# -*- coding: utf-8 -*-
"""build_docx.py — Báo cáo DOCX phân tích (10 mục như v4) từ view-model render.py.
Cột Nguồn là hyperlink; có ngày giá + độ tươi; mục Tin tức ở cuối. Đọc build_analysis (parity)."""
import datetime as dt
import pathlib
import sys

sys.path.append(str(pathlib.Path(__file__).resolve().parent.parent))
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from analytics import build_analysis
from report import render

OUT = pathlib.Path(__file__).resolve().parent.parent / "data"


def _add_hyperlink(paragraph, text, url):
    """Thêm hyperlink xanh gạch chân vào paragraph (python-docx không có sẵn)."""
    if not url:
        paragraph.add_run(text + " (form nội bộ)")
        return
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    hyper = OxmlElement("w:hyperlink"); hyper.set(qn("r:id"), r_id)
    run = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "2E6CA6"); rPr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    run.append(rPr)
    t = OxmlElement("w:t"); t.text = text; run.append(t)
    hyper.append(run); paragraph._p.append(hyper)


def _table(doc, headers, rows_data, best_flags=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for ri, row in enumerate(rows_data):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cells[ci].text = "" if val is None else str(val)
        if best_flags and best_flags[ri]:
            for c in cells:
                c.paragraphs[0].runs and setattr(c.paragraphs[0].runs[0].font, "bold", True)
    return t


def build(out_path=None, data=None):
    data = data or build_analysis.run()
    v = render.build_view(data)
    doc = Document()
    doc.add_heading("Báo cáo phân tích giá NVL Masterbatch — EUP Group", level=0)
    doc.add_paragraph(f"Cập nhật lúc {v['run_at']} · Phạm vi: nhựa nền + phụ gia (bỏ bột đá). "
                      f"Độ tươi 🟢≤7d 🟡≤30d 🔴>30d.")

    # ② KPI
    doc.add_heading("② Tổng quan nhanh (KPI)", level=1)
    rows = []
    for c in v["kpi"]:
        chg = "" if c["chg"] is None else (f"+{c['chg']:.1f}%" if c["chg"] > 0 else f"{c['chg']:.1f}%")
        rows.append([c["label"], c["value"], chg, c.get("landed") or "", c.get("date") or ""])
    _table(doc, ["NVL", "Giá", "Δ", "Landed", "Ngày giá"], rows)

    # ③ Phân tích
    doc.add_heading("③ Phân tích tổng hợp", level=1)
    for k in ("picture", "why", "impact", "reco"):
        doc.add_paragraph(v["narrative"][k])

    # ④ At-sight
    doc.add_heading("④ Quy đổi về GIÁ CHUNG — at-sight tương đương", level=1)
    for p, blk in v["atsight"].items():
        doc.add_heading(f"{p.upper()} — chênh tới {blk['gap']:,.0f} đ/kg", level=2)
        t = doc.add_table(rows=1, cols=8); t.style = "Light Grid Accent 1"
        for i, h in enumerate(["Khu vực", "Loại", "Thanh toán", "Giá gốc", "Landed", "Trả chậm", "At-sight", "Ngày"]):
            t.rows[0].cells[i].text = h
        for r in blk["rows"]:
            c = t.add_row().cells
            vals = [r["region"], r["price_type"], r["payment"], r["raw"],
                    f"{r['landed_vnd_kg']:,.0f}", f"{r['usance_benefit']:,.0f}",
                    f"{r['at_sight_equiv']:,.0f}", r["date"]]
            for i, val in enumerate(vals):
                c[i].text = str(val)
            if r["is_best"]:
                for cc in c:
                    if cc.paragraphs[0].runs:
                        cc.paragraphs[0].runs[0].font.bold = True

    # ⑤ Bảng tất cả khu vực (Nguồn là hyperlink)
    doc.add_heading("⑤ Cập nhật giá — tất cả khu vực", level=1)
    t = doc.add_table(rows=1, cols=7); t.style = "Light Grid Accent 1"
    for i, h in enumerate(["NVL", "Khu vực", "Loại", "Giá gốc", "USD/tấn", "Ngày", "Nguồn"]):
        t.rows[0].cells[i].text = h
    for r in v["allregion"]:
        c = t.add_row().cells
        for i, val in enumerate([r["product"], r["region"], r["price_type"], r["raw"],
                                 f"{r['usd_per_ton']:,.0f}", r["date"]]):
            c[i].text = str(val)
        _add_hyperlink(c[6].paragraphs[0], r["label"], r["url"])

    # ⑥ Insight chart (matplotlib PNG)
    doc.add_heading("⑥ Insight — chỉ số gốc-100", level=1)
    img = _index_chart_png(v)
    if img:
        doc.add_picture(str(img))

    # ⑦ Cảnh báo
    doc.add_heading("⑦ Cảnh báo & Đề xuất hành động", level=1)
    for a in v["alerts"]:
        pr = doc.add_paragraph()
        pr.add_run(f"[{a['severity']}] {a['title']}").bold = True
        doc.add_paragraph(a["detail"])
        doc.add_paragraph(f"→ Đề xuất: {a['recommendation']}")

    # ⑧ Dự báo
    doc.add_heading("⑧ Dự báo 6 tuần", level=1)
    for p, blk in v["forecast"].items():
        m = blk["meta"]
        doc.add_paragraph(f"{p.upper()} — model {m['model']}, Theil's U {m['theils_u']}, "
                          f"độ tin cậy {m['confidence']} ({m['basis']})")
        _table(doc, ["Tuần", "Dự báo", "Khoảng dưới", "Khoảng trên"],
               [[r["week"], f"{r['yhat']:,.0f}", f"{r['lower']:,.0f}", f"{r['upper']:,.0f}"]
                for r in blk["rows"]])

    # ⑨ Nguồn & độ tươi + Thuật ngữ
    doc.add_heading("⑨ Nguồn & độ tươi", level=1)
    t = doc.add_table(rows=1, cols=5); t.style = "Light Grid Accent 1"
    for i, h in enumerate(["Nguồn", "NVL", "Khu vực", "Ngày giá", "Tươi"]):
        t.rows[0].cells[i].text = h
    for s in v["sources"]:
        c = t.add_row().cells
        _add_hyperlink(c[0].paragraphs[0], s["label"], s["url"])
        c[1].text = s["product"]; c[2].text = s["region"]; c[3].text = s["date"]
        c[4].text = f"{s['fresh']['emoji']} {s['fresh']['days']}d"
    doc.add_heading("Thuật ngữ", level=2)
    for g in v["glossary"]:
        pr = doc.add_paragraph(); pr.add_run(g["term"] + ": ").bold = True; pr.add_run(g["desc"])

    # ⑩ Tin tức (cuối)
    doc.add_heading("⑩ Tin tức mới cập nhật", level=1)
    for it in v["news"]:
        pr = doc.add_paragraph(f"{it['published_at']} — ")
        _add_hyperlink(pr, it["title"], it["url"])
        pr.add_run(f"  ({it['source']})")
        doc.add_paragraph(it.get("summary", ""))

    out = pathlib.Path(out_path) if out_path else (OUT / f"bao_cao_phan_tich_{dt.date.today().isoformat()}.docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"[OK] report DOCX -> {out}")
    return str(out)


def _index_chart_png(v):
    """Vẽ chỉ số gốc-100 bằng matplotlib -> PNG tạm. Lỗi -> None (không gãy)."""
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from analytics import spreads as S
        fig, ax = plt.subplots(figsize=(6.5, 3.2))
        for p, ser in v["series"].items():
            if p == "brent":
                continue
            ys = S.index_100([val for _, val in ser])
            ax.plot(range(len(ys)), ys, label=p.upper())
        ax.axhline(100, ls=":", c="#888"); ax.legend(fontsize=7, ncol=4); ax.set_title("Chỉ số gốc-100")
        img = OUT / "_chart_index.png"; OUT.mkdir(parents=True, exist_ok=True)
        fig.tight_layout(); fig.savefig(img, dpi=110); plt.close(fig)
        return img
    except Exception as e:  # noqa: BLE001
        print(f"[WARN] chart docx lỗi: {e}")
        return None


if __name__ == "__main__":
    build()
