"""Sinh báo cáo phân tích dạng DOCX từ kết quả pipeline."""
from __future__ import annotations

import os
from datetime import datetime

from docx import Document
from docx.shared import Pt

from .models import RunResult


def _fmt(value) -> str:
    if value is None:
        return "—"
    if isinstance(value, float):
        return f"{value:,.2f}"
    if isinstance(value, int):
        return f"{value:,}"
    return str(value)


def build_report(result: RunResult, out_dir: str) -> str:
    os.makedirs(out_dir, exist_ok=True)
    doc = Document()

    title = doc.add_heading(f"Báo cáo phân tích cổ phiếu {result.ticker}", level=0)
    title.runs[0].font.size = Pt(20)

    ts = datetime.fromisoformat(result.generated_at).strftime("%Y-%m-%d %H:%M UTC")
    doc.add_paragraph(f"Thời điểm tạo: {ts}")
    doc.add_paragraph(
        "Nguồn phân tích: " + ("Claude API" if result.ai_used else "Phân tích theo luật")
    )

    # --- Tổng quan giá ---
    doc.add_heading("1. Tổng quan giá", level=1)
    ind = result.indicators
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Light Grid Accent 1"
    overview = [
        ("Giá đóng cửa gần nhất", _fmt(ind.last_close)),
        ("Phiên trước", _fmt(ind.prev_close)),
        ("Thay đổi", f"{_fmt(ind.change)} ({_fmt(ind.change_pct)}%)"),
        ("Khối lượng phiên gần nhất", _fmt(ind.last_volume)),
    ]
    for label, value in overview:
        row = tbl.add_row().cells
        row[0].text = label
        row[1].text = value

    # --- Chỉ báo kỹ thuật ---
    doc.add_heading("2. Chỉ báo kỹ thuật", level=1)
    tbl2 = doc.add_table(rows=0, cols=2)
    tbl2.style = "Light Grid Accent 1"
    indicators = [
        ("SMA(5)", _fmt(ind.sma5)),
        ("SMA(20)", _fmt(ind.sma20)),
        ("RSI(14)", _fmt(ind.rsi14)),
        ("Cao nhất kỳ", _fmt(ind.high_period)),
        ("Thấp nhất kỳ", _fmt(ind.low_period)),
        ("KL trung bình 20 phiên", _fmt(ind.avg_vol20)),
    ]
    for label, value in indicators:
        row = tbl2.add_row().cells
        row[0].text = label
        row[1].text = value

    # --- Nhận định ---
    doc.add_heading("3. Nhận định phân tích", level=1)
    for para in result.analysis.split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    # --- Dữ liệu 10 phiên gần nhất ---
    doc.add_heading("4. Dữ liệu 10 phiên gần nhất", level=1)
    tbl3 = doc.add_table(rows=1, cols=6)
    tbl3.style = "Light List Accent 1"
    hdr = tbl3.rows[0].cells
    for i, h in enumerate(["Ngày", "Mở", "Cao", "Thấp", "Đóng", "KL"]):
        hdr[i].text = h
    for bar in result.snapshot.bars[-10:]:
        cells = tbl3.add_row().cells
        cells[0].text = bar.date
        cells[1].text = _fmt(bar.open)
        cells[2].text = _fmt(bar.high)
        cells[3].text = _fmt(bar.low)
        cells[4].text = _fmt(bar.close)
        cells[5].text = _fmt(bar.volume)

    doc.add_paragraph(
        "\nMiễn trừ trách nhiệm: Báo cáo mang tính tham khảo, không phải khuyến "
        "nghị đầu tư."
    )

    stamp = datetime.fromisoformat(result.generated_at).strftime("%Y%m%d_%H%M%S")
    filename = f"{result.ticker}_{stamp}.docx"
    path = os.path.join(out_dir, filename)
    doc.save(path)
    return path
